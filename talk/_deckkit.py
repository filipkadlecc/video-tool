"""Shared slide furniture for the talk decks.

Brand tokens are the same ones the tool enforces on its own output (lib/brand.ts),
so these sit next to the Apify-template slides without a reskin. Both build scripts
import from here so the two decks can't drift apart.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

BG      = RGBColor(0x16, 0x17, 0x18)
BORDER  = RGBColor(0x3d, 0x3f, 0x43)
TEXT    = RGBColor(0xf4, 0xf4, 0xf5)
MUTED   = RGBColor(0xbf, 0xc1, 0xc5)
SUBTLE  = RGBColor(0x8c, 0x93, 0xa8)
ORANGE  = RGBColor(0xF8, 0x66, 0x06)

W, H = Inches(13.333), Inches(7.5)
M    = Inches(0.78)          # page margin

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]

def slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(1, 0, 0, W, H)           # 1 = rectangle
    bg.fill.solid(); bg.fill.fore_color.rgb = BG
    bg.line.fill.background(); bg.shadow.inherit = False
    return s

def rect(s, x, y, w, h, color):
    r = s.shapes.add_shape(1, x, y, w, h)
    r.fill.solid(); r.fill.fore_color.rgb = color
    r.line.fill.background(); r.shadow.inherit = False
    return r

def text(s, x, y, w, h, runs, size=18, color=TEXT, bold=False,
         spacing=1.15, align=PP_ALIGN.LEFT, caps=False, font="Inter"):
    """runs: a string, or [(text, {bold,color,size}), ...] for mixed styling."""
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = spacing
    if isinstance(runs, str):
        runs = [(runs, {})]
    for t, o in runs:
        r = p.add_run(); r.text = t
        f = r.font
        f.name = o.get("font", font)
        f.size = Pt(o.get("size", size))
        f.bold = o.get("bold", bold)
        f.color.rgb = o.get("color", color)
        if caps or o.get("caps"):
            r.text = t.upper()
    return tb

def head(s, tag, title_runs):
    """Orange rule + slide tag + headline — the same furniture as the HTML deck."""
    text(s, W - M - Inches(4.2), M - Inches(0.06), Inches(4.2), Inches(0.3),
         tag, size=10.5, color=SUBTLE, align=PP_ALIGN.RIGHT, caps=True, font="Roboto Mono")
    rect(s, M, M + Inches(0.12), Inches(0.72), Emu(28575), ORANGE)   # ~3px rule
    text(s, M, M + Inches(0.42), W - 2*M, Inches(1.1), title_runs,
         size=40, bold=True, spacing=1.0)

def points(s, rows, top=Inches(2.55), row_h=Inches(0.86), label_w=Inches(1.55), size=17):
    """Hairline-ruled list: small label, then the line, then the softer half."""
    y = top
    rect(s, M, y - Inches(0.14), W - 2*M, Emu(9525), BORDER)
    for label, strong, soft in rows:
        text(s, M, y + Inches(0.07), label_w, Inches(0.3), label,
             size=10, color=SUBTLE, caps=True, font="Roboto Mono")
        runs = [(strong, {"color": TEXT, "bold": False})]
        if soft:
            runs.append((" " + soft, {"color": MUTED}))
        text(s, M + label_w, y, W - 2*M - label_w, row_h, runs, size=size, spacing=1.3)
        y += row_h
        rect(s, M, y - Inches(0.14), W - 2*M, Emu(9525), BORDER)
    return y

def kicker(s, runs, y):
    text(s, M, y + Inches(0.3), W - 2*M - Inches(1.5), Inches(0.9), runs, size=21, bold=True, spacing=1.25)

